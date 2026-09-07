"""Build downloadable reports (PDF / Excel / CSV) from live database data.

Maintenance, Damage and Financial reports are *mapped* views of data the
system actually stores (no maintenance log, damage log or pricing exists):
- maintenance → low-stock items + overdue borrows needing attention
- damage      → late returns (returned after due date) as incident proxy
- financial   → stock composition per category (no cost fields available)
"""

import csv
import io
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path

from fpdf import FPDF, FontFace
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.models.borrow import BorrowRecord, BorrowStatus
from app.models.item import Item
from app.models.location import Location
from app.models.user import User, UserRoleEnum
from app.schemas.reports import ReportFormatEnum, ReportTypeEnum

FONT_DIR = Path(__file__).resolve().parent.parent / "assets" / "fonts"
DEFAULT_PERIOD_DAYS = 30
USAGE_TOP_LIMIT = 20


@dataclass
class ReportData:
    """Neutral shape a report builder produces before rendering."""

    title: str
    headers: list[str]
    rows: list[list]
    note: str | None = None
    period_start: date | None = None
    period_end: date | None = None


# ── Range helpers ─────────────────────────────────────────────────────────


def _resolve_period(
    start_date: date | None, end_date: date | None
) -> tuple[date, date]:
    """Default to the last N days when the caller omits the range."""
    end = end_date or datetime.now(timezone.utc).date()
    start = start_date or (end - timedelta(days=DEFAULT_PERIOD_DAYS))
    if start > end:
        start, end = end, start
    return start, end


def _utc_bounds(start: date, end: date) -> tuple[datetime, datetime]:
    """Inclusive day bounds as naive UTC datetimes (columns store UTC wall time)."""
    start_dt = datetime.combine(start, time.min, tzinfo=timezone.utc).replace(tzinfo=None)
    end_dt = datetime.combine(end, time.max, tzinfo=timezone.utc).replace(tzinfo=None)
    return start_dt, end_dt


def _fmt(value: datetime | None) -> str:
    return value.strftime("%Y-%m-%d %H:%M") if value else "-"


# ── Report builders (one per report type) ─────────────────────────────────


def _inventory_data(db: Session) -> ReportData:
    """Snapshot of every active item with quantities and stock status."""
    rows_db = (
        db.query(Item, Location.name)
        .outerjoin(Location, Item.location_id == Location.id)
        .filter(Item.is_active.is_(True))
        .order_by(Item.name)
        .all()
    )
    rows = [
        [
            item.id,
            item.name,
            item.category or "-",
            location_name or "-",
            item.total_quantity,
            item.available_quantity,
            item.total_quantity - item.available_quantity,
            "Low" if item.available_quantity <= item.low_stock_threshold else "OK",
        ]
        for item, location_name in rows_db
    ]
    return ReportData(
        title="Inventory Report",
        headers=[
            "ID", "Item", "Category", "Location", "Total Qty",
            "Available Qty", "Borrowed Qty", "Stock Status",
        ],
        rows=rows,
    )


def _borrowing_data(db: Session, user: User, start: date, end: date) -> ReportData:
    """Every borrow transaction inside the period (own records for non-admins)."""
    start_dt, end_dt = _utc_bounds(start, end)
    query = (
        db.query(BorrowRecord, User.full_name, Item.name)
        .join(User, BorrowRecord.user_id == User.id)
        .join(Item, BorrowRecord.item_id == Item.id)
        .filter(BorrowRecord.borrowed_at >= start_dt, BorrowRecord.borrowed_at <= end_dt)
    )
    if user.role != UserRoleEnum.admin:
        query = query.filter(BorrowRecord.user_id == user.id)
    records = query.order_by(BorrowRecord.borrowed_at.desc()).all()

    now = datetime.now(timezone.utc).replace(tzinfo=None)
    rows = [
        [
            record.id,
            borrower or "-",
            item_name,
            record.quantity,
            record.status.value,
            _fmt(record.borrowed_at),
            _fmt(record.due_date),
            _fmt(record.returned_at),
            _overdue_flag(record, now),
        ]
        for record, borrower, item_name in records
    ]
    return ReportData(
        title="Borrowing Report",
        headers=[
            "ID", "Borrower", "Item", "Qty", "Status",
            "Borrowed At", "Due Date", "Returned At", "Overdue",
        ],
        rows=rows,
        period_start=start,
        period_end=end,
    )


def _overdue_flag(record: BorrowRecord, now: datetime) -> str:
    if record.due_date is None:
        return "-"
    if record.status == BorrowStatus.returned:
        return "Yes" if record.returned_at and record.returned_at > record.due_date else "No"
    return "Yes" if record.due_date < now else "No"


def _usage_data(db: Session, user: User, start: date, end: date) -> ReportData:
    """Most-borrowed items inside the period (own usage for non-admins)."""
    start_dt, end_dt = _utc_bounds(start, end)
    query = (
        db.query(
            Item.name,
            Item.category,
            func.count(BorrowRecord.id).label("borrow_count"),
            func.coalesce(func.sum(BorrowRecord.quantity), 0).label("total_qty"),
        )
        .join(BorrowRecord, BorrowRecord.item_id == Item.id)
        .filter(BorrowRecord.borrowed_at >= start_dt, BorrowRecord.borrowed_at <= end_dt)
    )
    if user.role != UserRoleEnum.admin:
        query = query.filter(BorrowRecord.user_id == user.id)
    rows_db = (
        query.group_by(Item.id, Item.name, Item.category)
        .order_by(func.count(BorrowRecord.id).desc())
        .limit(USAGE_TOP_LIMIT)
        .all()
    )
    rows = [
        [rank, name, category or "-", int(borrow_count), int(total_qty)]
        for rank, (name, category, borrow_count, total_qty) in enumerate(rows_db, start=1)
    ]
    return ReportData(
        title="Usage Analytics Report",
        headers=["Rank", "Item", "Category", "Borrow Count", "Total Qty Borrowed"],
        rows=rows,
        note=f"Top {USAGE_TOP_LIMIT} items by borrow transactions.",
        period_start=start,
        period_end=end,
    )


def _maintenance_data(db: Session) -> ReportData:
    """Items needing attention: low stock + borrows currently overdue."""
    low_stock = (
        db.query(Item)
        .filter(Item.is_active.is_(True), Item.available_quantity <= Item.low_stock_threshold)
        .order_by(Item.name)
        .all()
    )
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    overdue = (
        db.query(BorrowRecord, User.full_name, Item.name)
        .join(User, BorrowRecord.user_id == User.id)
        .join(Item, BorrowRecord.item_id == Item.id)
        .filter(
            BorrowRecord.status == BorrowStatus.borrowed,
            BorrowRecord.due_date.isnot(None),
            BorrowRecord.due_date < now,
        )
        .order_by(BorrowRecord.due_date)
        .all()
    )
    rows = [
        ["Low Stock", item.id, item.name, f"Available {item.available_quantity} <= threshold {item.low_stock_threshold}"]
        for item in low_stock
    ] + [
        ["Overdue Borrow", record.id, item_name, f"Borrower: {borrower or '-'} · due {_fmt(record.due_date)}"]
        for record, borrower, item_name in overdue
    ]
    return ReportData(
        title="Maintenance Report",
        headers=["Type", "Ref ID", "Name", "Detail"],
        rows=rows,
        note="Mapped report: low-stock items and overdue borrows that need attention.",
    )


def _damage_data(db: Session, user: User, start: date, end: date) -> ReportData:
    """Late returns in the period as the closest proxy for damage incidents."""
    start_dt, end_dt = _utc_bounds(start, end)
    query = (
        db.query(BorrowRecord, User.full_name, Item.name)
        .join(User, BorrowRecord.user_id == User.id)
        .join(Item, BorrowRecord.item_id == Item.id)
        .filter(
            BorrowRecord.status == BorrowStatus.returned,
            BorrowRecord.due_date.isnot(None),
            BorrowRecord.returned_at.isnot(None),
            BorrowRecord.returned_at > BorrowRecord.due_date,
            BorrowRecord.returned_at >= start_dt,
            BorrowRecord.returned_at <= end_dt,
        )
    )
    if user.role != UserRoleEnum.admin:
        query = query.filter(BorrowRecord.user_id == user.id)
    records = query.order_by(BorrowRecord.returned_at.desc()).all()

    rows = [
        [
            record.id,
            borrower or "-",
            item_name,
            record.quantity,
            _fmt(record.due_date),
            _fmt(record.returned_at),
            _days_late(record),
        ]
        for record, borrower, item_name in records
    ]
    return ReportData(
        title="Damage Report",
        headers=["Borrow ID", "Borrower", "Item", "Qty", "Due Date", "Returned At", "Days Late"],
        rows=rows,
        note="Mapped report: late returns used as the incident proxy (no damage log in system).",
        period_start=start,
        period_end=end,
    )


def _days_late(record: BorrowRecord) -> int:
    if not record.due_date or not record.returned_at:
        return 0
    return max(0, (record.returned_at - record.due_date).days)


def _financial_data(db: Session) -> ReportData:
    """Stock composition per category — the closest available financial view."""
    rows_db = (
        db.query(
            func.coalesce(Item.category, "Uncategorized").label("category"),
            func.count(Item.id).label("item_types"),
            func.coalesce(func.sum(Item.total_quantity), 0).label("total_units"),
            func.coalesce(func.sum(Item.available_quantity), 0).label("available_units"),
        )
        .filter(Item.is_active.is_(True))
        .group_by(func.coalesce(Item.category, "Uncategorized"))
        .order_by(func.sum(Item.total_quantity).desc())
        .all()
    )
    rows = []
    for category, item_types, total_units, available_units in rows_db:
        borrowed_units = int(total_units) - int(available_units)
        utilization = round(borrowed_units / int(total_units) * 100, 1) if int(total_units) else 0.0
        rows.append(
            [category, int(item_types), int(total_units), int(available_units), borrowed_units, f"{utilization}%"]
        )
    return ReportData(
        title="Financial Summary Report",
        headers=["Category", "Item Types", "Total Units", "Available Units", "Borrowed Units", "Utilization"],
        rows=rows,
        note="Mapped report: stock composition per category (no pricing data in system).",
    )


_BUILDERS = {
    ReportTypeEnum.inventory: lambda db, user, s, e: _inventory_data(db),
    ReportTypeEnum.borrowing: _borrowing_data,
    ReportTypeEnum.maintenance: lambda db, user, s, e: _maintenance_data(db),
    ReportTypeEnum.damage: _damage_data,
    ReportTypeEnum.usage: _usage_data,
    ReportTypeEnum.financial: lambda db, user, s, e: _financial_data(db),
}


# ── Renderers ─────────────────────────────────────────────────────────────


def _meta_lines(data: ReportData, generated_at: str) -> list[str]:
    lines = [f"Generated: {generated_at}"]
    if data.period_start and data.period_end:
        lines.append(f"Period: {data.period_start} to {data.period_end}")
    if data.note:
        lines.append(data.note)
    return lines


def _render_csv(data: ReportData, generated_at: str) -> bytes:
    buf = io.StringIO()
    buf.write("\ufeff")  # BOM so Excel detects UTF-8 and renders Thai correctly
    writer = csv.writer(buf)
    writer.writerow([data.title])
    for line in _meta_lines(data, generated_at):
        writer.writerow([line])
    writer.writerow([])
    writer.writerow(data.headers)
    writer.writerows(data.rows)
    return buf.getvalue().encode("utf-8")


def _render_excel(data: ReportData, generated_at: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append([data.title])
    for line in _meta_lines(data, generated_at):
        ws.append([line])
    ws.append([])
    ws.append(data.headers)

    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color="4472C4")
    for row in data.rows:
        ws.append(row)

    for idx, header in enumerate(data.headers, start=1):
        values = [len(str(header))] + [len(str(row[idx - 1])) for row in data.rows]
        ws.column_dimensions[get_column_letter(idx)].width = min(max(values) + 2, 50)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class _ReportPDF(FPDF):
    def footer(self):
        self.set_y(-15)
        self.set_font("Sarabun", "", 8)
        self.cell(0, 10, f"Page {self.page_no()}", align="C")


def _render_pdf(data: ReportData, generated_at: str) -> bytes:
    pdf = _ReportPDF(orientation="L" if len(data.headers) >= 6 else "P")
    pdf.add_font("Sarabun", "", str(FONT_DIR / "Sarabun-Regular.ttf"))
    pdf.add_font("Sarabun", "B", str(FONT_DIR / "Sarabun-Bold.ttf"))
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()

    pdf.set_font("Sarabun", "B", 16)
    pdf.cell(0, 10, data.title, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Sarabun", "", 9)
    for line in _meta_lines(data, generated_at):
        pdf.cell(0, 5, line, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    pdf.set_font("Sarabun", "", 9)
    pdf.set_draw_color(200, 200, 200)
    heading = FontFace(emphasis="BOLD", fill_color=(68, 114, 196), color=(255, 255, 255))
    with pdf.table(
        line_height=5.5, padding=1.5, text_align="LEFT", headings_style=heading
    ) as table:
        head = table.row()
        for header in data.headers:
            head.cell(header)
        for row in data.rows:
            table_row = table.row()
            for value in row:
                table_row.cell(str(value))

    return bytes(pdf.output())


# ── Public API ────────────────────────────────────────────────────────────


def build_report(
    db: Session,
    user: User,
    report_type: ReportTypeEnum,
    fmt: ReportFormatEnum,
    start_date: date | None = None,
    end_date: date | None = None,
) -> tuple[str, str, bytes]:
    """Render a report → (filename, media_type, file bytes)."""
    start, end = _resolve_period(start_date, end_date)
    data = _BUILDERS[report_type](db, user, start, end)
    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    renderers = {
        ReportFormatEnum.csv: _render_csv,
        ReportFormatEnum.excel: _render_excel,
        ReportFormatEnum.pdf: _render_pdf,
    }
    content = renderers[fmt](data, generated_at)

    media_types = {
        ReportFormatEnum.csv: "text/csv; charset=utf-8",
        ReportFormatEnum.excel: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ReportFormatEnum.pdf: "application/pdf",
    }
    filename = f"{report_type.value}-report-{end.isoformat()}.{fmt.value}"
    return filename, media_types[fmt], content
